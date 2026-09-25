# services/stock_investigation_export.py
from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from direction.constants import SEUIL_ATTENTION_JOURS
from direction.services.stock_investigation_service import StockInvestigationService

# Le libellé suit direction.constants.SEUIL_ATTENTION_JOURS plutôt qu'un texte
# figé : export et affichage écran doivent toujours refléter le même seuil.
HEADERS = ["Superviseur", "Agent", f"Produits en circulation (> {SEUIL_ATTENTION_JOURS} jours)"]


def _formater_produits(produits):
    """Une ligne lisible par produit : quantité, nom, date de sortie — sans
    fournisseur ni valorisation (checklist terrain, pas un document comptable,
    décision mdmaiga 17/09/2026)."""
    return "\n".join(
        f"{p['quantite']} {p['produit_nom']} — depuis le {p['date_reference']:%d/%m/%Y} "
        f"({p['jours_ecoules']} j)"
        for p in produits
    )


class StockInvestigationExportService:
    """Export de la liste « Produits à investiguer » (sprint-14, révision
    17/09/2026) — regroupé par superviseur puis agent, format à remettre à un
    agent de vérification terrain, qui la parcourt et confirme la présence
    physique de chaque produit chez chaque agent listé."""

    @staticmethod
    def _groupes(lignes_annotees):
        return StockInvestigationService.regrouper_par_superviseur(lignes_annotees)

    @staticmethod
    def export_excel(lignes_annotees):
        groupes = StockInvestigationExportService._groupes(lignes_annotees)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Produits à investiguer"

        ws.append(HEADERS)
        for col in range(1, len(HEADERS) + 1):
            ws.cell(row=1, column=col).font = Font(bold=True)

        row_index = 2
        for groupe in groupes:
            nom_superviseur = groupe["superviseur"].full_name if groupe["superviseur"] else "—"
            for bloc_agent in groupe["agents"]:
                ws.cell(row=row_index, column=1, value=nom_superviseur)
                ws.cell(row=row_index, column=2, value=bloc_agent["agent"].full_name)
                cell_produits = ws.cell(
                    row=row_index, column=3,
                    value=_formater_produits(bloc_agent["produits"]),
                )
                cell_produits.alignment = Alignment(wrap_text=True, vertical="top")
                row_index += 1

        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["B"].width = 22
        ws.column_dimensions["C"].width = 60

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    @staticmethod
    def export_pdf(lignes_annotees):
        groupes = StockInvestigationExportService._groupes(lignes_annotees)

        buffer = BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=landscape(A4),
            title="Produits à investiguer",
        )

        styles = getSampleStyleSheet()
        cell_style = styles["BodyText"]
        elements = [
            Paragraph(
                "<b>Produits à investiguer — checklist terrain</b>", styles["Title"]
            ),
            Paragraph(
                f"Produits en circulation depuis plus de {SEUIL_ATTENTION_JOURS} jours sans vente enregistrée "
                ,
                styles["Normal"],
            ),
            Spacer(1, 12),
        ]

        for groupe in groupes:
            nom_superviseur = groupe["superviseur"].full_name if groupe["superviseur"] else "—"
            elements.append(Paragraph(f"<b>Superviseur : {nom_superviseur}</b>", styles["Heading3"]))

            data = [["Agent", HEADERS[2]]]
            for bloc_agent in groupe["agents"]:
                data.append([
                    bloc_agent["agent"].full_name,
                    Paragraph(_formater_produits(bloc_agent["produits"]).replace("\n", "<br/>"), cell_style),
                ])

            table = Table(data, repeatRows=1, colWidths=[150, 600])
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 10),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.lightgrey]),
            ]))
            elements.append(table)
            elements.append(Spacer(1, 16))

        doc.build(elements)

        buffer.seek(0)
        return buffer
