# services/agent_detail_export.py
from io import BytesIO

import openpyxl
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from core.models import DetailDistribution, Vente
from direction.services.agent_detail_service import AgentDetailService

HEADERS_DISTRIBUTIONS = ["Date", "Produit", "Quantité", "Superviseur", "Type distribution"]
HEADERS_VENTES = ["Date", "Produit", "Quantité", "Prix unitaire", "Montant", "Type vente", "Mode paiement"]
HEADERS_POSSESSION = ["Produit", "Quantité", "Remis le", "Depuis (j)"]


class AgentDetailExportService:
    """Export ligne à ligne des distributions reçues et des ventes réalisées par un agent,
    sur la période sélectionnée dans direction.AgentDetailView (demande mdmaiga 24/09/2026) —
    même filtre période (hebdo/mensuel/personnalisée) que la page, résolu par
    AgentDetailService.resolve_period."""

    @staticmethod
    def distributions_recues(agent, date_debut, date_fin):
        return (
            DetailDistribution.objects
            .filter(
                distribution__agent_terrain=agent,
                distribution__date_distribution__date__range=(date_debut, date_fin),
            )
            .select_related("distribution__superviseur__user", "lot__produit")
            .order_by("distribution__date_distribution")
        )

    @staticmethod
    def ventes_realisees(agent, date_debut, date_fin):
        return (
            Vente.objects
            .filter(
                agent=agent,
                date_vente__date__range=(date_debut, date_fin),
                est_supprime=False,
            )
            .select_related("detail_distribution__lot__produit")
            .order_by("date_vente")
        )

    @staticmethod
    def _ligne_distribution(d):
        return [
            d.distribution.date_distribution.strftime("%d/%m/%Y %H:%M"),
            d.lot.produit.nom,
            float(d.quantite),
            d.distribution.superviseur.full_name if d.distribution.superviseur else "—",
            d.distribution.get_type_distribution_display(),
        ]

    @staticmethod
    def _ligne_vente(v):
        return [
            v.date_vente.strftime("%d/%m/%Y %H:%M"),
            v.detail_distribution.lot.produit.nom,
            float(v.quantite),
            float(v.prix_vente_unitaire),
            float(v.quantite * v.prix_vente_unitaire),
            v.get_type_vente_display(),
            v.get_mode_paiement_display(),
        ]

    @staticmethod
    def _ligne_possession(p):
        return [
            p["produit_nom"],
            float(p["quantite"]),
            p["date_remise"].strftime("%d/%m/%Y"),
            p["jours_ecoules"],
        ]

    @staticmethod
    def export_excel(agent, date_debut, date_fin):
        distributions = AgentDetailExportService.distributions_recues(agent, date_debut, date_fin)
        ventes = AgentDetailExportService.ventes_realisees(agent, date_debut, date_fin)

        wb = openpyxl.Workbook()

        ws1 = wb.active
        ws1.title = "Distributions reçues"
        ws1.append(HEADERS_DISTRIBUTIONS)
        for col in range(1, len(HEADERS_DISTRIBUTIONS) + 1):
            ws1.cell(row=1, column=col).font = Font(bold=True)
        for d in distributions:
            ws1.append(AgentDetailExportService._ligne_distribution(d))
        for col, largeur in zip("ABCDE", (18, 24, 12, 22, 20)):
            ws1.column_dimensions[col].width = largeur

        ws2 = wb.create_sheet("Ventes réalisées")
        ws2.append(HEADERS_VENTES)
        for col in range(1, len(HEADERS_VENTES) + 1):
            ws2.cell(row=1, column=col).font = Font(bold=True)
        for v in ventes:
            ws2.append(AgentDetailExportService._ligne_vente(v))
        for col, largeur in zip("ABCDEFG", (18, 24, 12, 14, 14, 16, 16)):
            ws2.column_dimensions[col].width = largeur

        # Produits actuellement chez l'agent (reste > 0), pas bornés à la période
        # sélectionnée — état courant, à montrer à l'agent en personne (demande
        # mdmaiga, 25/09/2026).
        ws3 = wb.create_sheet("Produits en sa possession")
        ws3.append(HEADERS_POSSESSION)
        for col in range(1, len(HEADERS_POSSESSION) + 1):
            ws3.cell(row=1, column=col).font = Font(bold=True)
        for p in AgentDetailService.get_produits_en_possession(agent):
            ws3.append(AgentDetailExportService._ligne_possession(p))
        for col, largeur in zip("ABCD", (24, 12, 14, 12)):
            ws3.column_dimensions[col].width = largeur

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    @staticmethod
    def export_excel_possession(agent):
        """Export dédié à la seule section « Produits en sa possession » — liste
        simple à remettre à l'agent en personne, sans les onglets distributions/
        ventes de l'export complet. Demande mdmaiga, 25/09/2026."""
        produits = AgentDetailService.get_produits_en_possession(agent)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Produits en sa possession"
        ws.append(HEADERS_POSSESSION)
        for col in range(1, len(HEADERS_POSSESSION) + 1):
            ws.cell(row=1, column=col).font = Font(bold=True)
        for p in produits:
            ws.append(AgentDetailExportService._ligne_possession(p))
        for col, largeur in zip("ABCD", (24, 12, 14, 12)):
            ws.column_dimensions[col].width = largeur

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    @staticmethod
    def export_pdf_possession(agent):
        produits = AgentDetailService.get_produits_en_possession(agent)

        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, title=f"Produits en sa possession — {agent.full_name}")
        styles = getSampleStyleSheet()

        elements = [
            Paragraph(f"<b>{agent.full_name}</b> — produits en sa possession", styles["Title"]),
            Paragraph(
                "Produits reçus non encore soldés.",
                styles["Normal"],
            ),
            Spacer(1, 12),
            AgentDetailExportService._table(
                HEADERS_POSSESSION,
                [AgentDetailExportService._ligne_possession(p) for p in produits],
            ),
        ]

        doc.build(elements)
        buffer.seek(0)
        return buffer

    @staticmethod
    def _table(headers, lignes):
        table = Table([headers] + lignes, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 10),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.lightgrey]),
        ]))
        return table

    @staticmethod
    def export_pdf(agent, date_debut, date_fin):
        distributions = AgentDetailExportService.distributions_recues(agent, date_debut, date_fin)
        ventes = AgentDetailExportService.ventes_realisees(agent, date_debut, date_fin)

        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), title=f"Détail agent — {agent.full_name}")
        styles = getSampleStyleSheet()

        elements = [
            Paragraph(f"<b>{agent.full_name}</b> — distributions reçues et ventes réalisées", styles["Title"]),
            Paragraph(f"Période : {date_debut:%d/%m/%Y} → {date_fin:%d/%m/%Y}", styles["Normal"]),
            Spacer(1, 12),
            Paragraph("Distributions reçues", styles["Heading2"]),
            AgentDetailExportService._table(
                HEADERS_DISTRIBUTIONS,
                [AgentDetailExportService._ligne_distribution(d) for d in distributions],
            ),
            Spacer(1, 20),
            Paragraph("Ventes réalisées", styles["Heading2"]),
            AgentDetailExportService._table(
                HEADERS_VENTES,
                [AgentDetailExportService._ligne_vente(v) for v in ventes],
            ),
            Spacer(1, 20),
            Paragraph("Produits en sa possession", styles["Heading2"]),
            AgentDetailExportService._table(
                HEADERS_POSSESSION,
                [
                    AgentDetailExportService._ligne_possession(p)
                    for p in AgentDetailService.get_produits_en_possession(agent)
                ],
            ),
        ]

        doc.build(elements)
        buffer.seek(0)
        return buffer
