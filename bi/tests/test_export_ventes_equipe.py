"""Export des ventes d'une équipe (bi/agents/equipe/<id>/export/) : quantités en unités + kg
(1 sac x poids_unitaire_kg), évolution mensuelle sur 3 mois avec variation."""
from datetime import datetime
from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest
from django.urls import reverse
from django.utils import timezone

from core.models import Agent, DetailDistribution, DistributionAgent, LotEntrepot, Produit, Vente


@pytest.fixture
def equipe(db, client, django_user_model):
    direction = django_user_model.objects.create_user(username="mdmaiga", password="x")
    Agent.objects.create(user=direction, type_agent="direction")
    client.force_login(direction)

    sup = Agent.objects.create(
        user=django_user_model.objects.create_user(username="sup1", password="x"),
        type_agent="entrepot",
    )
    agent = Agent.objects.create(
        user=django_user_model.objects.create_user(username="ag1", password="x"),
        type_agent="terrain", superviseur=sup,
    )
    oignon = Produit.objects.create(nom="oignon", poids_unitaire_kg=Decimal("25"))
    aloco = Produit.objects.create(nom="aloco")

    def vente(produit, quantite, annee, mois, jour=10):
        lot = LotEntrepot.objects.create(
            produit=produit, quantite_initiale=Decimal("500"), quantite_restante=Decimal("500"),
            prix_achat_unitaire=Decimal("100"),
        )
        dist = DistributionAgent.objects.create(
            superviseur=sup, agent_terrain=agent, quantite_totale=Decimal("500"),
        )
        detail = DetailDistribution.objects.create(
            distribution=dist, lot=lot, quantite=Decimal("500"),
        )
        v = Vente.objects.create(
            agent=agent, detail_distribution=detail, quantite=Decimal(quantite),
            prix_vente_unitaire=Decimal("10"), type_vente="detail",
        )
        Vente.objects.filter(pk=v.pk).update(
            date_vente=timezone.make_aware(datetime(annee, mois, jour, 12))
        )

    vente(oignon, "2", 2026, 9)    # 50 kg en septembre
    vente(oignon, "4", 2026, 10)   # 100 kg en octobre
    vente(aloco, "30", 2026, 10)   # 30 kg en octobre
    return sup


def _url(nom, sup, **params):
    base = reverse(f"bi:{nom}", args=[sup.id])
    return base + "?" + "&".join(f"{k}={v}" for k, v in params.items())


def test_excel_quantites_et_evolution(client, equipe):
    r = client.get(_url("export_ventes_equipe_excel", equipe, annee=2026, mois=10))
    assert r.status_code == 200
    wb = openpyxl.load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Synthèse produits"]

    lignes = [[c for c in row if c is not None] for row in wb["Synthèse produits"].iter_rows(values_only=True)]
    # 4 sacs d'oignon = 100 kg ; aloco vrac = 30 kg
    assert ["oignon", 4.0, 100.0, 1, 0] in lignes
    assert ["aloco", "—", 30.0, 1, 0] in lignes
    assert any(l and "4 sacs vendus" in str(l[0]) and "130 kg" in str(l[0]) for l in lignes)
    # évolution : septembre 50 kg, octobre 130 kg -> +80 kg (+160 %)
    assert ["Octobre 2026", 4.0, 130.0, "+80.0", "+160.0 %"] in lignes
    assert ["Septembre 2026", 2.0, 50.0, "+50.0", "—"] in lignes


def test_filtre_produit_restreint_aussi_l_evolution(client, equipe):
    r = client.get(_url("export_ventes_equipe_excel", equipe, annee=2026, mois=10,
                        produit=Produit.objects.get(nom="aloco").id))
    wb = openpyxl.load_workbook(BytesIO(r.content))
    lignes = [[c for c in row if c is not None] for row in wb["Synthèse produits"].iter_rows(values_only=True)]
    assert ["Octobre 2026", 0.0, 30.0, "+30.0", "—"] in lignes
    assert not any(l and l[0] == "oignon" for l in lignes)


def test_pdf(client, equipe):
    r = client.get(_url("export_ventes_equipe_pdf", equipe, annee=2026, mois=10))
    assert r.status_code == 200
    assert r.content.startswith(b"%PDF")


def test_acces_refuse_hors_direction(client, equipe, django_user_model):
    client.logout()
    r = client.get(_url("export_ventes_equipe_pdf", equipe, annee=2026, mois=10))
    assert r.status_code == 302
