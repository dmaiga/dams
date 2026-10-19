# services/ventes_equipe_export.py
from collections import OrderedDict
from datetime import date
from decimal import Decimal
from io import BytesIO

import openpyxl
from django.db.models import DecimalField, F, Q, Sum
from django.db.models.functions import Coalesce, TruncMonth
from openpyxl.styles import Alignment, Font
from reportlab.platypus import Paragraph, Spacer

from core.models import Vente
from core.pdf_compact import nouveau_document, styles_compacts, tableau

MOIS_FR = [
    "", "janvier", "février", "mars", "avril", "mai", "juin",
    "juillet", "août", "septembre", "octobre", "novembre", "décembre",
]
NB_MOIS_EVOLUTION = 3

KG_EXPR = F("quantite") * Coalesce(
    F("detail_distribution__lot__produit__poids_unitaire_kg"), Decimal("1")
)


def _mois_moins_n(annee, mois, n):
    total = (annee * 12 + (mois - 1)) - n
    return total // 12, total % 12 + 1


def _kg(vente):
    """Même règle que Vente.quantite_en_kg : unités x poids unitaire du produit
    (ex. 1 sac d'oignon x 25 kg), ou quantité telle quelle pour un produit vrac."""
    poids = vente.detail_distribution.lot.produit.poids_unitaire_kg
    return vente.quantite * poids if poids else vente.quantite


class VentesEquipeExportService:
    """Export des quantités vendues par l'équipe d'un superviseur (hiérarchie ACTUELLE,
    Agent.superviseur_id — comme la fiche bi/agents/equipe/<id>/) sur la période filtrée.
    Performance du superviseur, sans CA/prix : sacs vendus (produits conditionnés), volume en kg
    (quantité x Produit.poids_unitaire_kg, ou quantité telle quelle pour le vrac), agents ayant
    vendu, évolution mensuelle sur 3 mois."""

    @staticmethod
    def collecter(superviseur, date_debut, date_fin, annee, mois, produit_id=None):
        # Équipe + ventes réalisées par le superviseur lui-même (colonne « Superviseur »).
        base = Vente.objects.filter(
            Q(agent__superviseur_id=superviseur.id) | Q(agent_id=superviseur.id),
            est_supprime=False,
        )
        if produit_id:
            base = base.filter(detail_distribution__lot__produit_id=produit_id)

        # Évolution mensuelle : NB_MOIS_EVOLUTION mois finissant au mois de la période,
        # + 1 mois avant servant uniquement de base à la variation du premier.
        annee_deb, mois_deb = _mois_moins_n(annee, mois, NB_MOIS_EVOLUTION)
        premier_jour = date(annee_deb, mois_deb, 1)
        par_mois = {}
        for r in (
            base.filter(date_vente__date__gte=premier_jour)
            .annotate(m=TruncMonth("date_vente"))
            .values("m")
            .annotate(
                kg=Sum(KG_EXPR, output_field=DecimalField(max_digits=15, decimal_places=2)),
                sacs=Sum("quantite", filter=Q(detail_distribution__lot__produit__poids_unitaire_kg__isnull=False)),
            )
        ):
            m = r["m"]
            par_mois[date(m.year, m.month, 1)] = (r["kg"], r["sacs"] or Decimal("0"))

        evolution = []
        precedent = par_mois.get(premier_jour, (Decimal("0"), Decimal("0")))[0]
        for n in range(NB_MOIS_EVOLUTION - 1, -1, -1):
            a, m = _mois_moins_n(annee, mois, n)
            kg, sacs = par_mois.get(date(a, m, 1), (Decimal("0"), Decimal("0")))
            variation = kg - precedent
            pct = float(variation / precedent * 100) if precedent else None
            evolution.append({
                "libelle": f"{MOIS_FR[m].capitalize()} {a}",
                "kg": kg,
                "sacs": sacs,
                "variation": variation,
                "variation_pct": pct,
            })
            precedent = kg

        # Ventes de la période
        ventes = list(
            base.filter(date_vente__date__range=(date_debut, date_fin))
            .select_related("agent__user", "detail_distribution__lot__produit")
            .order_by("date_vente", "agent__user__last_name")
        )

        produits = OrderedDict()
        agents_total = set()
        for v in ventes:
            produit = v.detail_distribution.lot.produit
            conditionne = bool(produit.poids_unitaire_kg)
            p = produits.setdefault(produit.nom, {
                "conditionne": conditionne, "sacs": Decimal("0"),
                "kg": Decimal("0"), "agents": set(), "superviseur": False,
            })
            if conditionne:
                p["sacs"] += v.quantite
            p["kg"] += _kg(v)
            if v.agent_id == superviseur.id:
                p["superviseur"] = True
            else:
                p["agents"].add(v.agent_id)
                agents_total.add(v.agent_id)

        produits_tries = sorted(produits.items(), key=lambda x: x[0])
        return {
            "superviseur_nom": superviseur.full_name,
            "date_debut": date_debut,
            "date_fin": date_fin,
            "produit_nom": produits_tries[0][0] if produit_id and produits_tries else None,
            "evolution": evolution,
            "produits": produits_tries,
            "total_sacs": sum((p["sacs"] for p in produits.values()), Decimal("0")),
            "total_kg": sum((p["kg"] for p in produits.values()), Decimal("0")),
            "nb_agents": len(agents_total),
        }

    # ------------------------------------------------------------------ tableaux
    @staticmethod
    def _entete(d):
        produit = d["produit_nom"] or "tous"
        return [
            f"Ventes de l'équipe — Superviseur : {d['superviseur_nom']}",
            f"Période : {d['date_debut']:%d/%m/%Y} → {d['date_fin']:%d/%m/%Y}   |   Produit : {produit}",
            f"Performance du superviseur : {float(d['total_sacs']):g} sacs vendus · "
            f"{d['total_kg']:.0f} kg (volume) · {d['nb_agents']} agent(s) ayant vendu",
        ]

    @staticmethod
    def _lignes_evolution(d):
        lignes = [["Mois", "Sacs vendus", "Kg vendus", "Variation (kg)", "Variation (%)"]]
        for e in d["evolution"]:
            lignes.append([
                e["libelle"], float(e["sacs"]), round(float(e["kg"]), 1),
                f"{float(e['variation']):+.1f}",
                f"{e['variation_pct']:+.1f} %" if e["variation_pct"] is not None else "—",
            ])
        return lignes

    @staticmethod
    def _lignes_produits(d):
        lignes = [["Produit", "Sacs vendus", "Quantité (kg)", "Agents ayant vendu", "Superviseur"]]
        for nom, p in d["produits"]:
            lignes.append([
                nom,
                float(p["sacs"]) if p["conditionne"] else "—",
                round(float(p["kg"]), 1),
                len(p["agents"]),
                1 if p["superviseur"] else 0,
            ])
        return lignes

    # ------------------------------------------------------------------ Excel
    @classmethod
    def export_excel(cls, d):
        wb = openpyxl.Workbook()

        def bloc(ws, titre, lignes):
            ws.append([titre])
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
            for i, ligne in enumerate(lignes):
                ws.append(ligne)
                if i == 0:
                    for c in range(1, len(ligne) + 1):
                        ws.cell(row=ws.max_row, column=c).font = Font(bold=True)

        def entete(ws):
            for i, texte in enumerate(cls._entete(d)):
                ws.append([texte])
                if i == 0:
                    ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)
            ws.append([])

        ws = wb.active
        ws.title = "Synthèse produits"
        entete(ws)
        bloc(ws, "Évolution mensuelle (sacs et kg vendus)", cls._lignes_evolution(d))
        ws.append([])
        bloc(ws, "Synthèse par produit", cls._lignes_produits(d))

        for feuille in wb.worksheets:
            for lettre, largeur in zip("ABCDE", (26, 16, 16, 16, 20)):
                feuille.column_dimensions[lettre].width = largeur
            for ligne in feuille.iter_rows():
                for cellule in ligne:
                    # Textes d'introduction : une seule cellule A, sans retour à la ligne
                    # forcé ; en-têtes de tableaux centrés et passés à la ligne si besoin.
                    if isinstance(cellule.font, Font) and cellule.font.b and cellule.column > 1:
                        cellule.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
                    elif cellule.column > 1:
                        cellule.alignment = Alignment(horizontal="center")

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    # ------------------------------------------------------------------ PDF
    @classmethod
    def export_pdf(cls, d):
        buffer = BytesIO()
        doc = nouveau_document(buffer, "Ventes de l'équipe")
        styles = styles_compacts()
        info = cls._entete(d)
        elements = [
            Paragraph(f"<b>{info[0]}</b>", styles["Title"]),
            Paragraph(info[1], styles["Normal"]),
            Paragraph(info[2], styles["Normal"]),
            Spacer(1, 6),
            Paragraph("Évolution mensuelle (sacs et kg vendus)", styles["Heading2"]),
            tableau(cls._lignes_evolution(d), doc, fractions=[4, 2, 2, 2, 2]),
            Paragraph("Synthèse par produit", styles["Heading2"]),
            tableau(cls._lignes_produits(d), doc, fractions=[4, 2, 2, 2, 2]),
        ]
        doc.build(elements)
        buffer.seek(0)
        return buffer
